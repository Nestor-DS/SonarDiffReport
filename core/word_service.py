from __future__ import annotations

import logging
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from PIL import Image

from core.errors import WordError
from models.comparison import ComparisonEntry
from models.file_change import ChangeStatus

logger = logging.getLogger(__name__)

_COL_ARCHIVO_CM = 5.0
_COL_IMAGEN_CM = 9.5
_DPI_REFERENCE = 96  # ancho natural de la imagen a 96 DPI


class WordService:
    """Genera el .docx con la tabla comparativa Antes/Después."""

    def generate(
        self,
        output_path: Path,
        base_branch: str,
        compare_branch: str,
        entries: list[ComparisonEntry],
        repo_path: str,
    ) -> Path:
        try:
            doc = Document()
            self._configure_page(doc)
            self._add_header(doc, base_branch, compare_branch, repo_path)
            self._add_table(doc, entries)
            output_path.parent.mkdir(parents=True, exist_ok=True)
            doc.save(str(output_path))
        except OSError as exc:
            raise WordError(f"No se pudo escribir el documento en {output_path}: {exc}") from exc
        logger.info("Documento generado: %s (%d entradas)", output_path, len(entries))
        return output_path

    @staticmethod
    def _configure_page(doc: Document) -> None:
        # Landscape para que quepan las dos capturas lado a lado
        section = doc.sections[0]
        width, height = section.page_width, section.page_height
        if width < height:
            section.page_width, section.page_height = height, width
        normal = doc.styles["Normal"]
        normal.font.size = Pt(10)

    def _add_header(self, doc: Document, base: str, compare: str, repo_path: str) -> None:
        title = doc.add_heading("Comparación de issues — SonarQube", level=1)
        title.alignment = WD_ALIGN_PARAGRAPH.LEFT

        fields = [
            ("Rama base (Antes): ", base),
            ("Rama a comparar (Después): ", compare),
            ("Repositorio: ", repo_path),
            ("Fecha: ", date.today().strftime("%d/%m/%Y")),
        ]
        for label, value in fields:
            paragraph = doc.add_paragraph()
            paragraph.add_run(label).bold = True
            paragraph.add_run(value)
        doc.add_paragraph()

    def _add_table(self, doc: Document, entries: list[ComparisonEntry]) -> None:
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"

        headers = ("Archivo", "Antes", "Después")
        for cell, text in zip(table.rows[0].cells, headers):
            cell.text = ""
            run = cell.paragraphs[0].add_run(text)
            run.bold = True

        widths = (Cm(_COL_ARCHIVO_CM), Cm(_COL_IMAGEN_CM), Cm(_COL_IMAGEN_CM))
        for cell, width in zip(table.rows[0].cells, widths):
            cell.width = width

        for entry in entries:
            row = table.add_row()
            for cell, width in zip(row.cells, widths):
                cell.width = width
            self._fill_row(row, entry)

    def _fill_row(self, row: object, entry: ComparisonEntry) -> None:
        cells = row.cells  # type: ignore[attr-defined]
        change = entry.change

        # Columna Archivo
        cells[0].text = ""
        p = cells[0].paragraphs[0]
        p.add_run(change.label).bold = True
        status_text = {
            ChangeStatus.ADDED: "(nuevo)",
            ChangeStatus.DELETED: "(eliminado)",
            ChangeStatus.RENAMED: "(renombrado)",
        }.get(change.status)
        if status_text:
            cells[0].add_paragraph(status_text)

        # Columna Antes
        if entry.before_image is not None:
            self._insert_image(cells[1], entry.before_image)
        elif change.status is ChangeStatus.ADDED:
            cells[1].text = "— (archivo nuevo)"
        else:
            cells[1].text = "— (sin captura)"

        # Columna Después
        if entry.after_image is not None:
            self._insert_image(cells[2], entry.after_image)
        elif change.status is ChangeStatus.DELETED:
            cells[2].text = "— (archivo eliminado)"
        else:
            cells[2].text = "— (sin captura)"

    @staticmethod
    def _insert_image(cell: object, image_path: Path) -> None:
        cell.text = ""  # type: ignore[attr-defined]
        paragraph = cell.paragraphs[0]  # type: ignore[attr-defined]
        try:
            with Image.open(image_path) as image:
                pixel_width, _ = image.size
            natural_cm = pixel_width * 2.54 / _DPI_REFERENCE
            width = min(Cm(natural_cm), Cm(_COL_IMAGEN_CM))
            paragraph.add_run().add_picture(str(image_path), width=width)
        except OSError as exc:
            logger.warning("No se pudo insertar la imagen %s: %s", image_path, exc)
            paragraph.add_run("— (error al cargar imagen)")
